"""Hospital admin visit-slot allotment: quota, routines, even-split materialization."""

from __future__ import annotations

import io
import json
import logging
from datetime import date, datetime, time, timedelta
from typing import Any

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from sqlalchemy import func
from sqlalchemy.orm import Session, joinedload

from app.models import (
    Branch,
    BranchVisitSlotPolicy,
    DoctorAvailabilitySlot,
    User,
    VisitSlotAllotment,
    VisitSlotAllotmentSource,
    VisitSlotRoutine,
)
from app.models.enums import Role
from app.services.messaging_service import EmailService
from app.utils.timezone import now_ist

logger = logging.getLogger(__name__)

DEFAULT_DAILY_QUOTA = 50
MANAGE_ROLES = {
    Role.SUPER_ADMIN.value,
    Role.HOSPITAL_ADMIN.value,
    Role.BRANCH_ADMIN.value,
}

TEMPLATE_HEADERS = (
    "Date",
    "StartTime",
    "EndTime",
    "SlotCount",
    "StaffName",
    "StaffEmail",
)
HEADER_ALIASES = {
    "date": "Date",
    "starttime": "StartTime",
    "start": "StartTime",
    "start_time": "StartTime",
    "endtime": "EndTime",
    "end": "EndTime",
    "end_time": "EndTime",
    "slotcount": "SlotCount",
    "slots": "SlotCount",
    "slot_count": "SlotCount",
    "staffname": "StaffName",
    "name": "StaffName",
    "staff": "StaffName",
    "staff_name": "StaffName",
    "staffemail": "StaffEmail",
    "email": "StaffEmail",
    "staff_email": "StaffEmail",
}


def even_split_slots(
    day: date,
    window_start: str,
    window_end: str,
    slot_count: int,
) -> list[tuple[datetime, datetime]]:
    """Return (slotStart, slotEnd) pairs evenly filling [window_start, window_end)."""
    if slot_count < 1:
        raise ValueError("slot_count must be >= 1")
    sh, sm = _parse_hhmm(window_start)
    eh, em = _parse_hhmm(window_end)
    start = datetime(day.year, day.month, day.day, sh, sm, 0, 0)
    end = datetime(day.year, day.month, day.day, eh, em, 0, 0)
    if end <= start:
        raise ValueError("window_end must be after window_start")
    total_seconds = (end - start).total_seconds()
    step = total_seconds / slot_count
    if step < 60:
        raise ValueError("Each slot must be at least 1 minute long")
    out: list[tuple[datetime, datetime]] = []
    for i in range(slot_count):
        slot_start = start + timedelta(seconds=step * i)
        slot_end = start + timedelta(seconds=step * (i + 1))
        # Snap to whole minutes
        slot_start = slot_start.replace(second=0, microsecond=0)
        slot_end = slot_end.replace(second=0, microsecond=0)
        if slot_end <= slot_start:
            slot_end = slot_start + timedelta(minutes=1)
        out.append((slot_start, slot_end))
    return out


def _parse_hhmm(value: str) -> tuple[int, int]:
    raw = (value or "").strip()
    try:
        parts = raw.split(":")
        hour = int(parts[0])
        minute = int(parts[1]) if len(parts) > 1 else 0
        if hour < 0 or hour > 23 or minute < 0 or minute > 59:
            raise ValueError("out of range")
        return hour, minute
    except (ValueError, IndexError) as exc:
        raise ValueError("Invalid time. Use HH:MM.") from exc


def _parse_date(value: str | date | datetime | None) -> date:
    if value is None:
        return now_ist().date()
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return datetime.strptime(str(value).strip()[:10], "%Y-%m-%d").date()
    except ValueError as exc:
        raise HTTPException(status_code=400, detail="Invalid date. Use YYYY-MM-DD.") from exc


class VisitSlotAllotmentService:
    def __init__(self, db: Session) -> None:
        self.db = db

    def _assert_branch_access(self, user: dict, branch_id: str) -> None:
        role = user.get("role")
        if role == Role.SUPER_ADMIN.value:
            return
        if user.get("branchId") != branch_id:
            raise HTTPException(status_code=403, detail="You do not have access to this branch.")

    def _require_manage(self, user: dict) -> None:
        if user.get("role") not in MANAGE_ROLES:
            raise HTTPException(status_code=403, detail="Not allowed to manage visit slots.")

    def get_or_create_policy(self, branch_id: str) -> BranchVisitSlotPolicy:
        policy = (
            self.db.query(BranchVisitSlotPolicy)
            .filter(BranchVisitSlotPolicy.branchId == branch_id)
            .first()
        )
        if policy:
            return policy
        policy = BranchVisitSlotPolicy(branchId=branch_id, dailyQuota=DEFAULT_DAILY_QUOTA)
        self.db.add(policy)
        self.db.flush()
        return policy

    def get_policy(self, user: dict, branch_id: str) -> dict:
        self._assert_branch_access(user, branch_id)
        policy = self.get_or_create_policy(branch_id)
        self.db.commit()
        return {"branchId": branch_id, "dailyQuota": policy.dailyQuota}

    def update_policy(self, user: dict, branch_id: str, daily_quota: int) -> dict:
        self._require_manage(user)
        self._assert_branch_access(user, branch_id)
        if daily_quota < 1 or daily_quota > 5000:
            raise HTTPException(status_code=400, detail="dailyQuota must be between 1 and 5000.")
        policy = self.get_or_create_policy(branch_id)
        used = self._day_allotted_count(branch_id, now_ist().date())
        if daily_quota < used:
            raise HTTPException(
                status_code=400,
                detail=f"Cannot set quota below today's allotted slots ({used}).",
            )
        policy.dailyQuota = daily_quota
        self.db.commit()
        return {"branchId": branch_id, "dailyQuota": policy.dailyQuota}

    def list_allotable_staff(self, user: dict, branch_id: str) -> dict:
        self._assert_branch_access(user, branch_id)
        rows = (
            self.db.query(User)
            .filter(
                User.branchId == branch_id,
                User.role == Role.STAFF.value,
                User.isActive == True,  # noqa: E712
            )
            .order_by(User.name)
            .all()
        )
        return {
            "items": [
                {
                    "id": u.id,
                    "name": u.name,
                    "email": u.email,
                    "userType": u.userType,
                    "departmentId": u.departmentId,
                    "subDepartmentId": u.subDepartmentId,
                }
                for u in rows
            ]
        }

    def _assert_allotable_staff(self, branch_id: str, staff_id: str) -> User:
        staff = self.db.get(User, staff_id)
        if (
            not staff
            or not staff.isActive
            or staff.branchId != branch_id
            or staff.role != Role.STAFF.value
        ):
            raise HTTPException(status_code=404, detail="Staff not found at this branch.")
        return staff

    def _parse_window(self, start: str, end: str) -> tuple[str, str]:
        try:
            sh, sm = _parse_hhmm(start)
            eh, em = _parse_hhmm(end)
        except ValueError as exc:
            raise HTTPException(status_code=400, detail=str(exc)) from exc
        start_n = f"{sh:02d}:{sm:02d}"
        end_n = f"{eh:02d}:{em:02d}"
        if (eh, em) <= (sh, sm):
            raise HTTPException(status_code=400, detail="endTime must be after startTime.")
        return start_n, end_n

    def _day_allotted_count(
        self,
        branch_id: str,
        allotment_date: date,
        *,
        exclude_allotment_id: str | None = None,
    ) -> int:
        q = self.db.query(func.coalesce(func.sum(VisitSlotAllotment.slotCount), 0)).filter(
            VisitSlotAllotment.branchId == branch_id,
            VisitSlotAllotment.allotmentDate == allotment_date,
        )
        if exclude_allotment_id:
            q = q.filter(VisitSlotAllotment.id != exclude_allotment_id)
        return int(q.scalar() or 0)

    def _assert_quota(
        self,
        branch_id: str,
        allotment_date: date,
        additional: int,
        *,
        exclude_allotment_id: str | None = None,
    ) -> None:
        policy = self.get_or_create_policy(branch_id)
        used = self._day_allotted_count(
            branch_id, allotment_date, exclude_allotment_id=exclude_allotment_id
        )
        if used + additional > policy.dailyQuota:
            raise HTTPException(
                status_code=409,
                detail=(
                    f"Daily visit-slot quota exceeded "
                    f"({used + additional} > {policy.dailyQuota})."
                ),
            )

    def _window_bounds(self, day: date, window_start: str, window_end: str) -> tuple[datetime, datetime]:
        sh, sm = _parse_hhmm(window_start)
        eh, em = _parse_hhmm(window_end)
        start = datetime(day.year, day.month, day.day, sh, sm, 0, 0)
        end = datetime(day.year, day.month, day.day, eh, em, 0, 0)
        return start, end

    def materialize_allotment(self, allotment: VisitSlotAllotment) -> dict:
        """Regenerate unbooked DoctorAvailabilitySlot rows for this allotment window."""
        try:
            desired = even_split_slots(
                allotment.allotmentDate,
                allotment.windowStart,
                allotment.windowEnd,
                allotment.slotCount,
            )
        except ValueError as exc:
            raise HTTPException(status_code=400, detail=str(exc)) from exc

        win_start, win_end = self._window_bounds(
            allotment.allotmentDate, allotment.windowStart, allotment.windowEnd
        )
        existing = (
            self.db.query(DoctorAvailabilitySlot)
            .filter(
                DoctorAvailabilitySlot.doctorId == allotment.staffId,
                DoctorAvailabilitySlot.slotStart >= win_start,
                DoctorAvailabilitySlot.slotStart < win_end,
            )
            .all()
        )
        booked = [s for s in existing if s.isBooked]
        if len(booked) > allotment.slotCount:
            raise HTTPException(
                status_code=409,
                detail=(
                    f"Cannot reduce to {allotment.slotCount} slots; "
                    f"{len(booked)} are already booked."
                ),
            )

        desired_starts = {s[0] for s in desired}
        now = now_ist()

        for slot in existing:
            if slot.isBooked:
                continue
            if slot.slotStart not in desired_starts:
                self.db.delete(slot)

        existing_by_start = {
            s.slotStart: s
            for s in self.db.query(DoctorAvailabilitySlot)
            .filter(
                DoctorAvailabilitySlot.doctorId == allotment.staffId,
                DoctorAvailabilitySlot.slotStart >= win_start,
                DoctorAvailabilitySlot.slotStart < win_end,
            )
            .all()
        }

        created = 0
        skipped_past = 0
        for slot_start, slot_end in desired:
            if slot_start in existing_by_start:
                continue
            if slot_start <= now:
                skipped_past += 1
                continue
            self.db.add(
                DoctorAvailabilitySlot(
                    doctorId=allotment.staffId,
                    slotStart=slot_start,
                    slotEnd=slot_end,
                    isBooked=False,
                )
            )
            created += 1

        self.db.flush()
        return {"created": created, "skippedPast": skipped_past, "bookedKept": len(booked)}

    def _serialize_allotment(self, row: VisitSlotAllotment) -> dict:
        staff = row.staff
        win_start, win_end = self._window_bounds(row.allotmentDate, row.windowStart, row.windowEnd)
        slots = (
            self.db.query(DoctorAvailabilitySlot)
            .filter(
                DoctorAvailabilitySlot.doctorId == row.staffId,
                DoctorAvailabilitySlot.slotStart >= win_start,
                DoctorAvailabilitySlot.slotStart < win_end,
            )
            .order_by(DoctorAvailabilitySlot.slotStart)
            .all()
        )
        booked = sum(1 for s in slots if s.isBooked)
        open_count = len(slots) - booked
        return {
            "id": row.id,
            "branchId": row.branchId,
            "staffId": row.staffId,
            "staffName": staff.name if staff else None,
            "staffUserType": staff.userType if staff else None,
            "allotmentDate": row.allotmentDate.isoformat(),
            "windowStart": row.windowStart,
            "windowEnd": row.windowEnd,
            "slotCount": row.slotCount,
            "source": row.source,
            "booked": booked,
            "open": open_count,
            "previewTimes": [
                s.slotStart.strftime("%H:%M") for s in slots
            ],
        }

    def list_allotments(self, user: dict, branch_id: str, date_str: str | None) -> dict:
        self._assert_branch_access(user, branch_id)
        day = _parse_date(date_str)
        policy = self.get_or_create_policy(branch_id)
        rows = (
            self.db.query(VisitSlotAllotment)
            .options(joinedload(VisitSlotAllotment.staff))
            .filter(
                VisitSlotAllotment.branchId == branch_id,
                VisitSlotAllotment.allotmentDate == day,
            )
            .order_by(VisitSlotAllotment.windowStart, VisitSlotAllotment.staffId)
            .all()
        )
        used = self._day_allotted_count(branch_id, day)
        self.db.commit()
        return {
            "date": day.isoformat(),
            "dailyQuota": policy.dailyQuota,
            "used": used,
            "remaining": max(0, policy.dailyQuota - used),
            "items": [self._serialize_allotment(r) for r in rows],
        }

    def create_allotment(
        self,
        user: dict,
        branch_id: str,
        *,
        staff_id: str,
        allotment_date: str,
        window_start: str,
        window_end: str,
        slot_count: int,
        source: str = VisitSlotAllotmentSource.MANUAL.value,
    ) -> dict:
        self._require_manage(user)
        self._assert_branch_access(user, branch_id)
        if slot_count < 1 or slot_count > 500:
            raise HTTPException(status_code=400, detail="slotCount must be between 1 and 500.")
        self._assert_allotable_staff(branch_id, staff_id)
        day = _parse_date(allotment_date)
        start_n, end_n = self._parse_window(window_start, window_end)
        existing = (
            self.db.query(VisitSlotAllotment)
            .filter(
                VisitSlotAllotment.staffId == staff_id,
                VisitSlotAllotment.allotmentDate == day,
                VisitSlotAllotment.windowStart == start_n,
            )
            .first()
        )
        if existing:
            raise HTTPException(
                status_code=409,
                detail="Allotment already exists for this staff, date, and window start. Use PATCH.",
            )
        self._assert_quota(branch_id, day, slot_count)
        row = VisitSlotAllotment(
            branchId=branch_id,
            staffId=staff_id,
            allotmentDate=day,
            windowStart=start_n,
            windowEnd=end_n,
            slotCount=slot_count,
            source=source,
        )
        self.db.add(row)
        self.db.flush()
        self.materialize_allotment(row)
        self.db.commit()
        self.db.refresh(row)
        row = (
            self.db.query(VisitSlotAllotment)
            .options(joinedload(VisitSlotAllotment.staff))
            .filter(VisitSlotAllotment.id == row.id)
            .one()
        )
        return self._serialize_allotment(row)

    def update_allotment(
        self,
        user: dict,
        branch_id: str,
        allotment_id: str,
        *,
        window_start: str | None = None,
        window_end: str | None = None,
        slot_count: int | None = None,
    ) -> dict:
        self._require_manage(user)
        self._assert_branch_access(user, branch_id)
        row = self.db.get(VisitSlotAllotment, allotment_id)
        if not row or row.branchId != branch_id:
            raise HTTPException(status_code=404, detail="Allotment not found.")

        old_start, old_end = self._window_bounds(row.allotmentDate, row.windowStart, row.windowEnd)
        new_start_s = window_start or row.windowStart
        new_end_s = window_end or row.windowEnd
        start_n, end_n = self._parse_window(new_start_s, new_end_s)
        new_count = slot_count if slot_count is not None else row.slotCount
        if new_count < 1 or new_count > 500:
            raise HTTPException(status_code=400, detail="slotCount must be between 1 and 500.")

        self._assert_quota(branch_id, row.allotmentDate, new_count, exclude_allotment_id=row.id)

        # Clear unbooked slots from previous window if window moves
        if (start_n, end_n) != (row.windowStart, row.windowEnd):
            for slot in (
                self.db.query(DoctorAvailabilitySlot)
                .filter(
                    DoctorAvailabilitySlot.doctorId == row.staffId,
                    DoctorAvailabilitySlot.slotStart >= old_start,
                    DoctorAvailabilitySlot.slotStart < old_end,
                    DoctorAvailabilitySlot.isBooked == False,  # noqa: E712
                )
                .all()
            ):
                self.db.delete(slot)
            self.db.flush()

        row.windowStart = start_n
        row.windowEnd = end_n
        row.slotCount = new_count
        if row.source == VisitSlotAllotmentSource.ROUTINE.value:
            row.source = VisitSlotAllotmentSource.OVERRIDE.value
        elif row.source != VisitSlotAllotmentSource.MANUAL.value:
            row.source = VisitSlotAllotmentSource.OVERRIDE.value

        self.materialize_allotment(row)
        self.db.commit()
        row = (
            self.db.query(VisitSlotAllotment)
            .options(joinedload(VisitSlotAllotment.staff))
            .filter(VisitSlotAllotment.id == allotment_id)
            .one()
        )
        return self._serialize_allotment(row)

    def delete_allotment(self, user: dict, branch_id: str, allotment_id: str) -> dict:
        self._require_manage(user)
        self._assert_branch_access(user, branch_id)
        row = self.db.get(VisitSlotAllotment, allotment_id)
        if not row or row.branchId != branch_id:
            raise HTTPException(status_code=404, detail="Allotment not found.")
        win_start, win_end = self._window_bounds(row.allotmentDate, row.windowStart, row.windowEnd)
        booked = (
            self.db.query(DoctorAvailabilitySlot)
            .filter(
                DoctorAvailabilitySlot.doctorId == row.staffId,
                DoctorAvailabilitySlot.slotStart >= win_start,
                DoctorAvailabilitySlot.slotStart < win_end,
                DoctorAvailabilitySlot.isBooked == True,  # noqa: E712
            )
            .count()
        )
        if booked:
            raise HTTPException(
                status_code=409,
                detail="Cannot delete allotment with booked slots. Reject visits first.",
            )
        for slot in (
            self.db.query(DoctorAvailabilitySlot)
            .filter(
                DoctorAvailabilitySlot.doctorId == row.staffId,
                DoctorAvailabilitySlot.slotStart >= win_start,
                DoctorAvailabilitySlot.slotStart < win_end,
            )
            .all()
        ):
            self.db.delete(slot)
        self.db.delete(row)
        self.db.commit()
        return {"id": allotment_id, "deleted": True}

    def _serialize_routine(self, row: VisitSlotRoutine) -> dict:
        staff = row.staff
        try:
            weekdays = json.loads(row.weekdays or "[]")
        except json.JSONDecodeError:
            weekdays = []
        return {
            "id": row.id,
            "branchId": row.branchId,
            "staffId": row.staffId,
            "staffName": staff.name if staff else None,
            "weekdays": weekdays,
            "windowStart": row.windowStart,
            "windowEnd": row.windowEnd,
            "slotCount": row.slotCount,
            "isActive": bool(row.isActive),
        }

    def list_routines(self, user: dict, branch_id: str) -> dict:
        self._assert_branch_access(user, branch_id)
        rows = (
            self.db.query(VisitSlotRoutine)
            .options(joinedload(VisitSlotRoutine.staff))
            .filter(VisitSlotRoutine.branchId == branch_id)
            .order_by(VisitSlotRoutine.windowStart)
            .all()
        )
        return {"items": [self._serialize_routine(r) for r in rows]}

    def create_routine(
        self,
        user: dict,
        branch_id: str,
        *,
        staff_id: str,
        weekdays: list[int],
        window_start: str,
        window_end: str,
        slot_count: int,
    ) -> dict:
        self._require_manage(user)
        self._assert_branch_access(user, branch_id)
        self._assert_allotable_staff(branch_id, staff_id)
        if slot_count < 1 or slot_count > 500:
            raise HTTPException(status_code=400, detail="slotCount must be between 1 and 500.")
        for d in weekdays:
            if d < 0 or d > 6:
                raise HTTPException(status_code=400, detail="weekdays must be 0 (Mon) through 6 (Sun).")
        start_n, end_n = self._parse_window(window_start, window_end)
        row = VisitSlotRoutine(
            branchId=branch_id,
            staffId=staff_id,
            weekdays=json.dumps(sorted(set(weekdays))),
            windowStart=start_n,
            windowEnd=end_n,
            slotCount=slot_count,
            isActive=True,
        )
        self.db.add(row)
        self.db.commit()
        self.db.refresh(row)
        row = (
            self.db.query(VisitSlotRoutine)
            .options(joinedload(VisitSlotRoutine.staff))
            .filter(VisitSlotRoutine.id == row.id)
            .one()
        )
        return self._serialize_routine(row)

    def update_routine(
        self,
        user: dict,
        branch_id: str,
        routine_id: str,
        *,
        weekdays: list[int] | None = None,
        window_start: str | None = None,
        window_end: str | None = None,
        slot_count: int | None = None,
        is_active: bool | None = None,
    ) -> dict:
        self._require_manage(user)
        self._assert_branch_access(user, branch_id)
        row = self.db.get(VisitSlotRoutine, routine_id)
        if not row or row.branchId != branch_id:
            raise HTTPException(status_code=404, detail="Routine not found.")
        if weekdays is not None:
            for d in weekdays:
                if d < 0 or d > 6:
                    raise HTTPException(
                        status_code=400, detail="weekdays must be 0 (Mon) through 6 (Sun)."
                    )
            row.weekdays = json.dumps(sorted(set(weekdays)))
        if window_start is not None or window_end is not None:
            start_n, end_n = self._parse_window(
                window_start or row.windowStart, window_end or row.windowEnd
            )
            row.windowStart = start_n
            row.windowEnd = end_n
        if slot_count is not None:
            if slot_count < 1 or slot_count > 500:
                raise HTTPException(status_code=400, detail="slotCount must be between 1 and 500.")
            row.slotCount = slot_count
        if is_active is not None:
            row.isActive = is_active
        self.db.commit()
        row = (
            self.db.query(VisitSlotRoutine)
            .options(joinedload(VisitSlotRoutine.staff))
            .filter(VisitSlotRoutine.id == routine_id)
            .one()
        )
        return self._serialize_routine(row)

    def delete_routine(self, user: dict, branch_id: str, routine_id: str) -> dict:
        self._require_manage(user)
        self._assert_branch_access(user, branch_id)
        row = self.db.get(VisitSlotRoutine, routine_id)
        if not row or row.branchId != branch_id:
            raise HTTPException(status_code=404, detail="Routine not found.")
        self.db.delete(row)
        self.db.commit()
        return {"id": routine_id, "deleted": True}

    def apply_routines(
        self,
        user: dict,
        branch_id: str,
        *,
        from_date: str,
        to_date: str,
    ) -> dict:
        self._require_manage(user)
        self._assert_branch_access(user, branch_id)
        start = _parse_date(from_date)
        end = _parse_date(to_date)
        if end < start:
            raise HTTPException(status_code=400, detail="to date must be on or after from date.")
        if (end - start).days > 60:
            raise HTTPException(status_code=400, detail="Apply range cannot exceed 60 days.")

        routines = (
            self.db.query(VisitSlotRoutine)
            .filter(
                VisitSlotRoutine.branchId == branch_id,
                VisitSlotRoutine.isActive == True,  # noqa: E712
            )
            .all()
        )
        created = 0
        updated = 0
        skipped = 0
        cursor = start
        while cursor <= end:
            weekday = cursor.weekday()
            for routine in routines:
                try:
                    days = set(json.loads(routine.weekdays or "[]"))
                except json.JSONDecodeError:
                    days = set()
                if weekday not in days:
                    continue
                existing = (
                    self.db.query(VisitSlotAllotment)
                    .filter(
                        VisitSlotAllotment.staffId == routine.staffId,
                        VisitSlotAllotment.allotmentDate == cursor,
                        VisitSlotAllotment.windowStart == routine.windowStart,
                    )
                    .first()
                )
                if existing and existing.source in (
                    VisitSlotAllotmentSource.MANUAL.value,
                    VisitSlotAllotmentSource.OVERRIDE.value,
                ):
                    skipped += 1
                    continue
                if existing:
                    try:
                        self._assert_quota(
                            branch_id,
                            cursor,
                            routine.slotCount,
                            exclude_allotment_id=existing.id,
                        )
                    except HTTPException:
                        skipped += 1
                        continue
                    existing.windowEnd = routine.windowEnd
                    existing.slotCount = routine.slotCount
                    existing.source = VisitSlotAllotmentSource.ROUTINE.value
                    self.materialize_allotment(existing)
                    updated += 1
                else:
                    try:
                        self._assert_quota(branch_id, cursor, routine.slotCount)
                    except HTTPException:
                        skipped += 1
                        continue
                    row = VisitSlotAllotment(
                        branchId=branch_id,
                        staffId=routine.staffId,
                        allotmentDate=cursor,
                        windowStart=routine.windowStart,
                        windowEnd=routine.windowEnd,
                        slotCount=routine.slotCount,
                        source=VisitSlotAllotmentSource.ROUTINE.value,
                    )
                    self.db.add(row)
                    self.db.flush()
                    self.materialize_allotment(row)
                    created += 1
            cursor += timedelta(days=1)

        self.db.commit()
        return {"created": created, "updated": updated, "skipped": skipped}

    def build_template_xlsx(self) -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.title = "Allotments"
        ws.append(list(TEMPLATE_HEADERS))
        ws.append(["2026-08-25", "09:00", "10:00", 3, "AI Doctor Priya", "priya.nair@connitor-elcity.com"])
        ws.append(["2026-08-25", "10:00", "11:00", 2, "Lakshmi Devi", ""])
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def _cell_str(self, value: Any) -> str:
        if value is None:
            return ""
        if isinstance(value, datetime):
            return value.strftime("%Y-%m-%d")
        if isinstance(value, date):
            return value.isoformat()
        if isinstance(value, time):
            return value.strftime("%H:%M")
        if isinstance(value, float) and value == int(value):
            return str(int(value))
        return str(value).strip()

    def _parse_excel_date(self, value: Any) -> date:
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        raw = self._cell_str(value)
        if not raw:
            raise ValueError("Date is required")
        return _parse_date(raw)

    def _parse_excel_time(self, value: Any) -> str:
        if isinstance(value, datetime):
            return value.strftime("%H:%M")
        if isinstance(value, time):
            return value.strftime("%H:%M")
        if isinstance(value, (int, float)):
            # Excel serial time fraction of day
            total = int(round(float(value) * 24 * 60)) % (24 * 60)
            return f"{total // 60:02d}:{total % 60:02d}"
        raw = self._cell_str(value)
        if not raw:
            raise ValueError("Time is required")
        # Accept HH:MM or HH:MM:SS
        parts = raw.replace(".", ":").split(":")
        hour = int(parts[0])
        minute = int(parts[1]) if len(parts) > 1 else 0
        if hour < 0 or hour > 23 or minute < 0 or minute > 59:
            raise ValueError("Invalid time")
        return f"{hour:02d}:{minute:02d}"

    def _resolve_staff(
        self,
        branch_id: str,
        *,
        staff_name: str,
        staff_email: str,
    ) -> User:
        staff_rows = (
            self.db.query(User)
            .filter(
                User.branchId == branch_id,
                User.role == Role.STAFF.value,
                User.isActive == True,  # noqa: E712
            )
            .all()
        )
        email = staff_email.strip().lower()
        if email:
            matches = [u for u in staff_rows if (u.email or "").strip().lower() == email]
            if len(matches) == 1:
                return matches[0]
            if not matches:
                raise ValueError(f"No staff found with email {staff_email}")
            raise ValueError(f"Multiple staff match email {staff_email}")

        name = staff_name.strip().lower()
        if not name:
            raise ValueError("StaffName or StaffEmail is required")
        matches = [u for u in staff_rows if (u.name or "").strip().lower() == name]
        if len(matches) == 1:
            return matches[0]
        if not matches:
            raise ValueError(f"No staff found named '{staff_name}'")
        raise ValueError(f"Multiple staff named '{staff_name}'; add StaffEmail")

    def _notify_staff_allotment(self, branch_id: str, allotment: dict) -> None:
        staff = self.db.get(User, allotment["staffId"])
        if not staff or not staff.email:
            return
        branch = self.db.get(Branch, branch_id)
        branch_name = branch.name if branch else "your hospital"
        times = ", ".join(allotment.get("previewTimes") or []) or "see dashboard"
        subject = f"Visitor slots allotted — {allotment['allotmentDate']}"
        message = (
            f"Hello {staff.name or 'Staff'},\n\n"
            f"Visitor appointment slots have been allotted for you at {branch_name}.\n\n"
            f"Date: {allotment['allotmentDate']}\n"
            f"Window: {allotment['windowStart']} – {allotment['windowEnd']}\n"
            f"Slots: {allotment['slotCount']}\n"
            f"Times: {times}\n\n"
            f"Visitors can book these times for meetings with you.\n"
        )
        try:
            EmailService().send_notification(staff.email, subject, message)
        except Exception:
            logger.exception("Failed to email visit-slot allotment to %s", staff.email)

    def import_allotments_from_xlsx(self, user: dict, branch_id: str, file_bytes: bytes) -> dict:
        self._require_manage(user)
        self._assert_branch_access(user, branch_id)
        if not file_bytes:
            raise HTTPException(status_code=400, detail="Empty file.")

        try:
            wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
        except Exception as exc:
            raise HTTPException(status_code=400, detail="Invalid Excel file. Upload a .xlsx workbook.") from exc

        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        try:
            header_row = next(rows_iter)
        except StopIteration:
            raise HTTPException(status_code=400, detail="Spreadsheet has no header row.") from None

        col_map: dict[str, int] = {}
        for idx, cell in enumerate(header_row or []):
            key = HEADER_ALIASES.get(self._cell_str(cell).lower().replace(" ", ""))
            if key:
                col_map[key] = idx
        required = ["Date", "StartTime", "EndTime", "SlotCount"]
        missing = [h for h in required if h not in col_map]
        if missing:
            raise HTTPException(
                status_code=400,
                detail=f"Missing required columns: {', '.join(missing)}. Expected {', '.join(TEMPLATE_HEADERS)}.",
            )
        if "StaffName" not in col_map and "StaffEmail" not in col_map:
            raise HTTPException(status_code=400, detail="Provide StaffName and/or StaffEmail column.")

        created = 0
        updated = 0
        failed = 0
        results: list[dict] = []

        for excel_row_num, raw in enumerate(rows_iter, start=2):
            if raw is None or all(v is None or self._cell_str(v) == "" for v in raw):
                continue

            def col(name: str) -> Any:
                idx = col_map.get(name)
                if idx is None or idx >= len(raw):
                    return None
                return raw[idx]

            try:
                day = self._parse_excel_date(col("Date"))
                start_time = self._parse_excel_time(col("StartTime"))
                end_time = self._parse_excel_time(col("EndTime"))
                slot_raw = col("SlotCount")
                if isinstance(slot_raw, (int, float)):
                    slot_count = int(slot_raw)
                else:
                    slot_count = int(self._cell_str(slot_raw))
                staff_name = self._cell_str(col("StaffName")) if "StaffName" in col_map else ""
                staff_email = self._cell_str(col("StaffEmail")) if "StaffEmail" in col_map else ""
                staff = self._resolve_staff(branch_id, staff_name=staff_name, staff_email=staff_email)
                start_n, end_n = self._parse_window(start_time, end_time)

                existing = (
                    self.db.query(VisitSlotAllotment)
                    .filter(
                        VisitSlotAllotment.staffId == staff.id,
                        VisitSlotAllotment.allotmentDate == day,
                        VisitSlotAllotment.windowStart == start_n,
                    )
                    .first()
                )
                if existing:
                    allotment = self.update_allotment(
                        user,
                        branch_id,
                        existing.id,
                        window_start=start_n,
                        window_end=end_n,
                        slot_count=slot_count,
                    )
                    updated += 1
                    action = "updated"
                else:
                    allotment = self.create_allotment(
                        user,
                        branch_id,
                        staff_id=staff.id,
                        allotment_date=day.isoformat(),
                        window_start=start_n,
                        window_end=end_n,
                        slot_count=slot_count,
                        source=VisitSlotAllotmentSource.MANUAL.value,
                    )
                    created += 1
                    action = "created"

                self._notify_staff_allotment(branch_id, allotment)
                results.append(
                    {
                        "row": excel_row_num,
                        "ok": True,
                        "action": action,
                        "staffName": allotment.get("staffName"),
                        "date": allotment.get("allotmentDate"),
                        "windowStart": allotment.get("windowStart"),
                        "windowEnd": allotment.get("windowEnd"),
                        "slotCount": allotment.get("slotCount"),
                    }
                )
            except HTTPException as exc:
                failed += 1
                detail = exc.detail if isinstance(exc.detail, str) else str(exc.detail)
                results.append({"row": excel_row_num, "ok": False, "error": detail})
            except Exception as exc:
                failed += 1
                results.append({"row": excel_row_num, "ok": False, "error": str(exc)})

        return {
            "created": created,
            "updated": updated,
            "failed": failed,
            "rows": results,
        }
